#!/usr/bin/env python
"""Build the researcher's allocation summary — who is asked which question.

Produces ``outputs/evaluation/study_kit/03_ADMIN_PENELITI/RINGKASAN_ALOKASI.xlsx``:
one workbook answering "what does each participant actually get, and how does
the corpus split across them", derived entirely from the frozen assignment
(``data/derived/eval_case_assignments.json``), the sealed condition key, and the
distributed workbooks themselves — so the ordering shown here is the ordering
the participant sees, not a re-derivation of it.

Sheets:

    Ringkasan          headline counts + per-person workload
    Semua Kasus        all 50 corpus questions, one row each, with every role
    Konselor           the 48 answer slots (2 counselors x 24), in workbook order
    Psikolog Tahap 1   the 56 counselor-answer ratings, in rating order
    Psikolog Tahap 2   the 44 briefing ratings, in rating order
    Distribusi         the crosstabs (condition x counselor x risk, form, shared)

RESEARCHER ONLY. This file joins blind `R###` codes to condition and counselor,
which is exactly what the raters must not see; it lands in `03_ADMIN_PENELITI/`
beside `SEALED_p2_key.csv` and is distributed to nobody.

    .venv/bin/python scripts/build_allocation_summary.py
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from depression_rag.evaluation.deidentify import deidentify  # noqa: E402

KIT = ROOT / "outputs/evaluation/study_kit"
QUESTIONS = ROOT / "data/derived/questions_alodokter_sample50.jsonl"
ASSIGNMENTS = ROOT / "data/derived/eval_case_assignments.json"
SEALED = KIT / "03_ADMIN_PENELITI/SEALED_p2_key.csv"
OUT = KIT / "03_ADMIN_PENELITI/RINGKASAN_ALOKASI.xlsx"

# The four counselor workbooks, in the order each counselor works them. The
# tahap is the counterbalancing: counselor 1 starts DENGAN, counselor 2 TANPA.
COUNSELOR_BOOKS = [
    ("counselor_1", "chatbot", 1, "Konselor_1_DENGAN_alat_bantu.xlsx"),
    ("counselor_1", "no_chatbot", 2, "Konselor_1_TANPA_alat_bantu.xlsx"),
    ("counselor_2", "no_chatbot", 1, "Konselor_2_TANPA_alat_bantu.xlsx"),
    ("counselor_2", "chatbot", 2, "Konselor_2_DENGAN_alat_bantu.xlsx"),
]
PSYCH_BOOKS = [("psychologist_1", "Psikolog_1.xlsx"), ("psychologist_2", "Psikolog_2.xlsx")]

COND_ID = {"chatbot": "DENGAN alat bantu", "no_chatbot": "TANPA alat bantu"}
FORM_ID = {"full": "lengkap", "light": "ringkas"}

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
SHARED_FILL = PatternFill("solid", fgColor="FCE4D6")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(bold=True, size=13, color="1F3864")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# --- inputs ------------------------------------------------------------------
def load_inputs() -> tuple[dict, list[dict], list[dict]]:
    questions = {}
    for line in QUESTIONS.read_text(encoding="utf-8").splitlines():
        if line.strip():
            rec = json.loads(line)
            questions[rec["question_id"]] = rec
    cases = json.load(ASSIGNMENTS.open(encoding="utf-8"))["cases"]
    sealed = list(csv.DictReader(SEALED.open(encoding="utf-8")))
    return questions, cases, sealed


def read_orders() -> tuple[dict, dict, dict]:
    """Pull the worklists straight out of the distributed workbooks.

    The codes never move once frozen, so reading them back is a check as much as
    a lookup: a workbook rebuilt with a different order would show up here.
    """
    counselor: dict[tuple[str, str], list[str]] = {}
    for who, cond, _tahap, fname in COUNSELOR_BOOKS:
        ws = load_workbook(KIT / "01_KONSELOR" / fname, read_only=True)["Jawaban"]
        counselor[(who, cond)] = [
            r[1] for r in ws.iter_rows(min_row=3, values_only=True) if r[1]
        ]
    p2_order: dict[str, list[str]] = {}
    p1_order: dict[str, list[tuple[str, str]]] = {}
    for who, fname in PSYCH_BOOKS:
        wb = load_workbook(KIT / "02_PSIKOLOG" / fname, read_only=True)
        p2_order[who] = [
            r[1] for r in wb["Jawaban Konselor"].iter_rows(min_row=3, values_only=True) if r[1]
        ]
        p1_order[who] = [
            (r[1], r[3]) for r in wb["Materi Chatbot"].iter_rows(min_row=3, values_only=True) if r[1]
        ]
    return counselor, p2_order, p1_order


# --- sheet helpers -----------------------------------------------------------
def title_row(ws, text: str, note: str, ncols: int) -> None:
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    ws["A2"] = note
    ws["A2"].font = Font(italic=True, size=9, color="595959")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.row_dimensions[2].height = 28


def header(ws, row: int, cols: list[tuple[str, int]]) -> None:
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(cols))}{row}"


def body(ws, row: int, values: list, shade: bool = False) -> None:
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if shade:
            c.fill = SHARED_FILL


def snippet(questions: dict, qid: str, n: int = 90) -> str:
    text = deidentify(qid, questions[qid]["question_title"], strict=False)
    return text if len(text) <= n else text[: n - 1] + "…"


# --- sheets ------------------------------------------------------------------
def sheet_summary(wb, counts, cases, sealed, p1_order) -> None:
    ws = wb.create_sheet("Ringkasan")
    title_row(ws, "Allocation summary — who is asked which question", (
        "Derived from the frozen assignment (seed 20260628) and the workbooks as distributed. "
        "RESEARCHER ONLY: the Tahap-1 sheet joins blind R### codes to condition and counselor."), 6)
    ws["A3"] = ("⚠  NEVER give this file to a counselor or a psychologist — it breaks the blinding.")
    ws["A3"].fill, ws["A3"].font = WARN_FILL, Font(bold=True, color="833C00")
    ws.merge_cells("A3:F3")

    n_counselor = sum(1 for c in cases if c["in_counselor_sample"])
    n_counselor_risk = sum(1 for c in cases if c["in_counselor_sample"] and c["risky"])
    n_p1 = sum(1 for c in cases if c["in_p1"])
    both_roles = sum(1 for c in cases if c["in_counselor_sample"] and c["in_p1"])
    unused = sum(1 for c in cases if not c["in_counselor_sample"] and not c["in_p1"])
    shared_p2 = sum(1 for r in sealed if r["raters"] == "both")
    shared_p2_cases = len({r["question_id"] for r in sealed if r["raters"] == "both"})
    shared_p1 = sum(1 for c in cases if c["p1_rater"] == "both")

    rows = [
        ("The corpus", "", ""),
        ("Questions in the frozen corpus", 50, "data/derived/questions_alodokter_sample50.jsonl"),
        ("…used by the counselors", n_counselor,
         f"{n_counselor_risk} risiko + {n_counselor - n_counselor_risk} non-risiko; "
         "12 per condition per counselor"),
        ("…used for psychologist briefing rating", n_p1, "the P1 union across both raters"),
        ("…used in BOTH roles", both_roles, "same question answered by a counselor and rated as a briefing"),
        ("…not used by any human", unused, "corpus only — still scored by the automatic eval"),
        ("", "", ""),
        ("Per counselor (x2 counselors)", "", ""),
        ("Questions answered", 24, "12 DENGAN alat bantu + 12 TANPA alat bantu"),
        ("Written answers produced", 24, "150–350 words each, own prose"),
        ("Per-case rating rows", 12, "DENGAN block only, filled right after each answer"),
        ("Exit questionnaire", 1, "«Kuesioner Akhir», once at the end of the DENGAN block"),
        ("Workbooks handed over", 2, "one per condition; order is the counterbalancing"),
        ("Estimated time", "~4–6 h", "fine across 2 days"),
        ("", "", ""),
        ("Per psychologist (x2 psychologists)", "", ""),
        ("Items rated", 50, "28 counselor answers + 22 chatbot briefings"),
        ("Tahap 1 — counselor answers", 28, "blind R### codes, condition never shown"),
        ("Tahap 2 — chatbot briefings", 22, "11 lengkap (with faithfulness) + 11 ringkas"),
        ("Workbooks handed over", 1, "self-contained: 27 sheets incl. the 22 packets"),
        ("Estimated time", "~5–6 h", "best as two sittings of ≤3 h"),
        ("", "", ""),
        ("Overlap for inter-rater agreement", "", ""),
        ("Counselor answers rated by both", shared_p2,
         f"{shared_p2_cases} questions x both counselors, so both conditions of each; "
         f"{shared_p2 * 2} of the 56 Tahap-1 ratings"),
        ("Briefings rated by both", shared_p1,
         f"{shared_p1 * 2} of the 44 Tahap-2 ratings; all lengkap form, so faithfulness agreement is covered"),
        ("Shared items inside each rater's 50", shared_p2 + shared_p1,
         f"{shared_p2} answers + {shared_p1} briefings; "
         f"{(shared_p2 + shared_p1) * 2} rating events across the two raters"),
        ("", "", ""),
        ("Total human rating events", "", ""),
        ("Counselor answers written", 48, "2 counselors x 24"),
        ("Psychologist ratings", 100, "2 raters x 50 items"),
    ]
    r = 5
    for label, value, note in rows:
        if label and value == "" and note == "":
            c = ws.cell(row=r, column=1, value=label)
            c.font = Font(bold=True, size=11, color="1F3864")
            c.fill = SUB_FILL
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        elif label:
            ws.cell(row=r, column=1, value=label).alignment = Alignment(wrap_text=True)
            v = ws.cell(row=r, column=2, value=value)
            v.font = Font(bold=True)
            v.alignment = Alignment(horizontal="center")
            n = ws.cell(row=r, column=3, value=note)
            n.font = Font(size=9, color="595959")
            n.alignment = Alignment(wrap_text=True)
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="The other sheets").font = Font(bold=True, size=11, color="1F3864")
    ws.cell(row=r, column=1).fill = SUB_FILL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    for name, what in [
        ("Semua Kasus", "all 50 corpus questions, one row each — every role a question plays"),
        ("Konselor", "the 48 answer slots in the order each counselor meets them"),
        ("Psikolog Tahap 1", "the 56 counselor-answer ratings, in rating order (carries the sealed key)"),
        ("Psikolog Tahap 2", "the 44 briefing ratings, in rating order"),
        ("Distribusi", "the crosstabs — balance by condition, counselor, risk and form"),
    ]:
        r += 1
        ws.cell(row=r, column=1, value=name).font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=what)
        c.font = Font(size=9, color="595959")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)

    for col, width in zip("ABCDEF", (42, 12, 22, 18, 18, 18)):
        ws.column_dimensions[col].width = width


def sheet_all_cases(wb, questions, cases, sealed) -> None:
    ws = wb.create_sheet("Semua Kasus")
    title_row(ws, "All 50 corpus questions — every role each one plays", (
        "One row per question. A blank counselor/psychologist column means the question is not used "
        "by that group; those questions are still scored by the automatic evaluation."), 12)
    header(ws, 4, [
        ("Kode\nkasus", 8), ("question_id", 12), ("Blok", 6), ("Risiko", 8),
        ("Judul pertanyaan", 46),
        ("Konselor 1", 16), ("Konselor 2", 16),
        ("Kode respons\nK1", 12), ("Kode respons\nK2", 12),
        ("Penilai Tahap 1\n(K1 / K2)", 20),
        ("Penilai Tahap 2\n(materi chatbot)", 18), ("Bentuk\nform", 10),
    ])
    by_q = {(r["question_id"], r["counselor"]): r for r in sealed}
    r = 5
    for case in cases:
        qid = case["question_id"]
        cond = case["counselor_condition"] or {}
        k1, k2 = by_q.get((qid, "counselor_1")), by_q.get((qid, "counselor_2"))
        body(ws, r, [
            case["study_id"], qid, case["block"], "risiko" if case["risky"] else "non-risiko",
            snippet(questions, qid, 120),
            COND_ID.get(cond.get("counselor_1"), ""), COND_ID.get(cond.get("counselor_2"), ""),
            k1["response_id"] if k1 else "", k2["response_id"] if k2 else "",
            " / ".join(x for x in [
                _short(k1["raters"]) if k1 else "", _short(k2["raters"]) if k2 else ""] if x),
            _short(case["p1_rater"]) if case["p1_rater"] else "",
            FORM_ID.get(case["p1_form"], ""),
        ], shade=(case["p1_rater"] == "both" or (k1 and k1["raters"] == "both")))
        r += 1


def _short(rater: str) -> str:
    return {"psychologist_1": "P1", "psychologist_2": "P2", "both": "P1+P2"}.get(rater, rater)


def sheet_counselor(wb, questions, cases, sealed, counselor_order) -> None:
    ws = wb.create_sheet("Konselor")
    title_row(ws, "Counselor worklist — 48 answer slots", (
        "The order is the counterbalancing and is not negotiable: counselor 1 works DENGAN then TANPA, "
        "counselor 2 TANPA then DENGAN. Each counselor sees all 24 questions, each in one condition only; "
        "the two counselors see every question in opposite conditions."), 11)
    header(ws, 4, [
        ("Konselor", 11), ("Tahap", 7), ("Kondisi", 18), ("Berkas", 30),
        ("No\ndalam berkas", 9), ("Kode\nkasus", 8), ("question_id", 12), ("Risiko", 9),
        ("Judul pertanyaan", 44), ("Kode\nrespons", 9), ("Dinilai oleh", 12),
    ])
    case_by_q = {c["question_id"]: c for c in cases}
    q_by_study = {c["study_id"]: c["question_id"] for c in cases}
    sealed_by = {(r["question_id"], r["counselor"], r["condition"]): r for r in sealed}
    r = 5
    for who, cond, tahap, fname in COUNSELOR_BOOKS:
        for pos, study_id in enumerate(counselor_order[(who, cond)], start=1):
            qid = q_by_study[study_id]
            case = case_by_q[qid]
            key = sealed_by[(qid, who, cond)]
            body(ws, r, [
                who.replace("counselor_", "Konselor "), f"{tahap} dari 2", COND_ID[cond], fname,
                pos, study_id, qid, "risiko" if case["risky"] else "non-risiko",
                snippet(questions, qid, 110), key["response_id"], _short(key["raters"]),
            ], shade=(key["raters"] == "both"))
            r += 1


def sheet_psych_t1(wb, questions, cases, sealed, p2_order) -> None:
    ws = wb.create_sheet("Psikolog Tahap 1")
    title_row(ws, "Psychologist Tahap 1 — counselor answers (56 ratings, 28 each)", (
        "SEALED COLUMNS (shaded header): counselor and condition are the blinding key — they exist here "
        "for your allocation check only and must not reach a rater. Row order = the order in the "
        "workbook, which is also the drift covariate. Peach rows are the 8 cases rated by both."), 9)
    header(ws, 4, [
        ("Psikolog", 11), ("Urutan\nbaris", 8), ("Kode\nrespons", 10), ("Kode\nkasus", 8),
        ("Risiko", 10), ("Judul pertanyaan", 48),
        ("SEALED —\nkonselor", 12), ("SEALED —\nkondisi", 18), ("Dinilai\nbersama?", 11),
    ])
    for col in (7, 8):
        ws.cell(row=4, column=col).fill = PatternFill("solid", fgColor="843C0C")
    case_by_q = {c["question_id"]: c for c in cases}
    sealed_by_r = {r["response_id"]: r for r in sealed}
    r = 5
    for who, order in p2_order.items():
        for pos, rid in enumerate(order, start=1):
            key = sealed_by_r[rid]
            case = case_by_q[key["question_id"]]
            body(ws, r, [
                _short(who), pos, rid, case["study_id"],
                "risiko" if case["risky"] else "non-risiko",
                snippet(questions, key["question_id"], 120),
                key["counselor"].replace("counselor_", "Konselor "), COND_ID[key["condition"]],
                "ya — kedua penilai" if key["raters"] == "both" else "",
            ], shade=(key["raters"] == "both"))
            r += 1


def sheet_psych_t2(wb, questions, cases, p1_order) -> None:
    ws = wb.create_sheet("Psikolog Tahap 2")
    title_row(ws, "Psychologist Tahap 2 — chatbot briefings (44 ratings, 22 each)", (
        "Nothing here is blind: Tahap 2 rates the chatbot's own briefing, so the case code is shown to the "
        "rater. «lengkap» carries Section 3 (faithfulness), «ringkas» skips it. All 8 shared cases are "
        "lengkap, so faithfulness agreement is unaffected by the reduction from 18 to 14 full forms."), 8)
    header(ws, 4, [
        ("Psikolog", 11), ("Urutan\nbaris", 8), ("Kode\nkasus", 8), ("Bentuk form", 12),
        ("Risiko", 10), ("Judul pertanyaan", 48), ("Lembar paket", 16), ("Dinilai\nbersama?", 11),
    ])
    case_by_study = {c["study_id"]: c for c in cases}
    r = 5
    for who, order in p1_order.items():
        for pos, (study_id, form) in enumerate(order, start=1):
            case = case_by_study[study_id]
            body(ws, r, [
                _short(who), pos, study_id, form,
                "risiko" if case["risky"] else "non-risiko",
                snippet(questions, case["question_id"], 120),
                f"Paket {study_id}",
                "ya — kedua penilai" if case["p1_rater"] == "both" else "",
            ], shade=(case["p1_rater"] == "both"))
            r += 1


def sheet_distribution(wb, cases, sealed, p1_order) -> None:
    ws = wb.create_sheet("Distribusi")
    title_row(ws, "Distribution — how the questions spread across people and cells", (
        "Every cell below is counted from the same frozen assignment the workbooks were built from. "
        "The design balances on condition x counselor first, then on risk."), 7)

    r = 4

    def block(title: str, note: str, head: list[str], rows: list[list]) -> None:
        nonlocal r
        c = ws.cell(row=r, column=1, value=title)
        c.font = Font(bold=True, size=11, color="1F3864")
        c.fill = SUB_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        r += 1
        if note:
            n = ws.cell(row=r, column=1, value=note)
            n.font = Font(italic=True, size=9, color="595959")
            n.alignment = Alignment(wrap_text=True)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
            r += 1
        for i, name in enumerate(head, start=1):
            h = ws.cell(row=r, column=i, value=name)
            h.fill, h.font, h.border = HEAD_FILL, HEAD_FONT, BORDER
            h.alignment = Alignment(wrap_text=True, horizontal="center")
        r += 1
        for values in rows:
            for i, v in enumerate(values, start=1):
                cell = ws.cell(row=r, column=i, value=v)
                cell.border = BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top",
                                           horizontal="center" if i > 1 else "left")
            r += 1
        r += 1

    risky = {c["question_id"]: c["risky"] for c in cases}
    counselor_cases = [c for c in cases if c["in_counselor_sample"]]

    block("1 — Corpus, by role",
          "The counselor sample is deliberately half risk-flagged; the P1 union is risk-weighted instead, "
          "because briefing safety is what P1 is meant to detect.",
          ["Role", "Questions", "risiko", "non-risiko"],
          [["Full corpus", 50, sum(1 for c in cases if c["risky"]), sum(1 for c in cases if not c["risky"])],
           ["Answered by counselors", len(counselor_cases),
            sum(1 for c in counselor_cases if c["risky"]), sum(1 for c in counselor_cases if not c["risky"])],
           ["Rated as briefings (P1 union)", sum(1 for c in cases if c["in_p1"]),
            sum(1 for c in cases if c["in_p1"] and c["risky"]),
            sum(1 for c in cases if c["in_p1"] and not c["risky"])],
           ["Both roles", sum(1 for c in cases if c["in_p1"] and c["in_counselor_sample"]), "", ""],
           ["Neither (automatic eval only)",
            sum(1 for c in cases if not c["in_p1"] and not c["in_counselor_sample"]), "", ""]])

    cell = Counter((r_["counselor"], r_["condition"]) for r_ in sealed)
    cell_risk = Counter((r_["counselor"], r_["condition"]) for r_ in sealed if risky[r_["question_id"]])
    block("2 — Counselor answers, by counselor x condition",
          "The crossover: every question is answered twice, once in each condition, by different counselors.",
          ["Counselor", "DENGAN alat bantu", "TANPA alat bantu", "Total", "of which risiko"],
          [[who.replace("counselor_", "Konselor "), cell[(who, "chatbot")], cell[(who, "no_chatbot")],
            cell[(who, "chatbot")] + cell[(who, "no_chatbot")],
            cell_risk[(who, "chatbot")] + cell_risk[(who, "no_chatbot")]]
           for who in ("counselor_1", "counselor_2")]
          + [["Total", cell[("counselor_1", "chatbot")] + cell[("counselor_2", "chatbot")],
              cell[("counselor_1", "no_chatbot")] + cell[("counselor_2", "no_chatbot")],
              sum(cell.values()), sum(cell_risk.values())]])

    t1 = Counter()
    for row in sealed:
        for who in (["psychologist_1", "psychologist_2"] if row["raters"] == "both" else [row["raters"]]):
            t1[(who, row["condition"])] += 1
            t1[(who, row["counselor"])] += 1
            t1[who] += 1
    block("3 — Psychologist Tahap 1 (counselor answers), by rater",
          "Each rater sees both conditions and both counselors in roughly equal share — but never learns which.",
          ["Rater", "Items", "DENGAN", "TANPA", "from Konselor 1", "from Konselor 2", "shared with other rater"],
          [[_short(who), t1[who], t1[(who, "chatbot")], t1[(who, "no_chatbot")],
            t1[(who, "counselor_1")], t1[(who, "counselor_2")],
            sum(1 for x in sealed if x["raters"] == "both")]
           for who in ("psychologist_1", "psychologist_2")])

    case_by_study = {c["study_id"]: c for c in cases}
    block("4 — Psychologist Tahap 2 (chatbot briefings), by rater",
          "«lengkap» adds Section 3 (faithfulness against the 5 passages); «ringkas» skips it.",
          ["Rater", "Packets", "lengkap", "ringkas", "risiko", "non-risiko", "shared with other rater"],
          [[_short(who), len(order),
            sum(1 for _, f in order if f == "lengkap"), sum(1 for _, f in order if f == "ringkas"),
            sum(1 for s, _ in order if case_by_study[s]["risky"]),
            sum(1 for s, _ in order if not case_by_study[s]["risky"]),
            sum(1 for c in cases if c["p1_rater"] == "both")]
           for who, order in p1_order.items()])

    shared_p2 = sum(1 for x in sealed if x["raters"] == "both")
    shared_p1 = sum(1 for c in cases if c["p1_rater"] == "both")
    block("5 — Total workload per person",
          "What each participant is actually asked to produce, end to end.",
          ["Person", "Main task", "Additional", "Files received", "Est. time"],
          [["Konselor 1", "24 answers (12 DENGAN + 12 TANPA)",
            "12 per-case ratings + 1 exit questionnaire", "2 workbooks + chatbot link", "~4–6 h"],
           ["Konselor 2", "24 answers (12 TANPA + 12 DENGAN)",
            "12 per-case ratings + 1 exit questionnaire", "2 workbooks + chatbot link", "~4–6 h"],
           ["Psikolog 1", "50 items (28 answers + 22 briefings)",
            f"{shared_p2 + shared_p1} of them shared for agreement", "1 workbook (27 sheets)", "~5–6 h"],
           ["Psikolog 2", "50 items (28 answers + 22 briefings)",
            f"{shared_p2 + shared_p1} of them shared for agreement", "1 workbook (27 sheets)", "~5–6 h"],
           ["Peneliti (you)", "no rating", "builds packets, holds the sealed key, runs the analysis",
            "the whole kit", "—"]])

    for col, width in zip("ABCDEFG", (34, 26, 22, 22, 20, 20, 22)):
        ws.column_dimensions[col].width = width


# --- checks ------------------------------------------------------------------
def check(cases, sealed, counselor_order, p2_order, p1_order) -> None:
    """Fail loudly if the workbooks and the frozen assignment ever disagree."""
    q_by_study = {c["study_id"]: c["question_id"] for c in cases}
    for (who, cond), order in counselor_order.items():
        expected = {c["question_id"] for c in cases
                    if (c["counselor_condition"] or {}).get(who) == cond}
        got = {q_by_study[s] for s in order}
        if expected != got:
            raise SystemExit(f"{who}/{cond}: workbook questions do not match the frozen assignment")
    for who, order in p2_order.items():
        expected = {r["response_id"] for r in sealed
                    if r["raters"] in (who, "both")}
        if expected != set(order):
            raise SystemExit(f"{who}: Tahap-1 worklist does not match SEALED_p2_key.csv")
    for who, order in p1_order.items():
        expected = {c["study_id"] for c in cases if c["p1_rater"] in (who, "both")}
        if expected != {s for s, _ in order}:
            raise SystemExit(f"{who}: Tahap-2 worklist does not match the frozen assignment")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path, default=OUT)
    args = ap.parse_args()

    questions, cases, sealed = load_inputs()
    counselor_order, p2_order, p1_order = read_orders()
    check(cases, sealed, counselor_order, p2_order, p1_order)

    wb = Workbook()
    wb.remove(wb.active)
    sheet_summary(wb, None, cases, sealed, p1_order)
    sheet_all_cases(wb, questions, cases, sealed)
    sheet_counselor(wb, questions, cases, sealed, counselor_order)
    sheet_psych_t1(wb, questions, cases, sealed, p2_order)
    sheet_psych_t2(wb, questions, cases, p1_order)
    sheet_distribution(wb, cases, sealed, p1_order)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote {args.out.relative_to(ROOT)}  ({len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
