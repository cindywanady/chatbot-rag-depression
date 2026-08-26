#!/usr/bin/env python
"""Build the counselor answer workbooks — one workbook per CONDITION per counselor.

Produces four self-contained .xlsx files in
``outputs/evaluation/study_kit/01_KONSELOR/``::

    Konselor_1_TANPA_alat_bantu.xlsx     Konselor_1_DENGAN_alat_bantu.xlsx
    Konselor_2_TANPA_alat_bantu.xlsx     Konselor_2_DENGAN_alat_bantu.xlsx

Each workbook opens on a "Petunjuk" sheet that a counselor reads once, top to
bottom, and then never has to come back to: it carries the task, the answer
template, the rules, the safety instruction, and the meaning + scale of every
column they will fill. The DENGAN workbooks additionally carry the exit
questionnaire on a third sheet.

Design constraints implemented here (rationale and item provenance live in
``documents/evaluation/INTERNAL_catatan_item_konselor.md`` — a researcher-only
file, never handed to a counselor):

* every construct is measured at exactly ONE level (per case OR exit, never both);
* nothing in the counselor-facing sheets names a source instrument, an internal
  item code, the source corpus, or the internal risk flag;
* the patient questions are name-scrubbed before they are shown.

    .venv/bin/python scripts/build_counselor_workbooks.py --archive-old
"""
from __future__ import annotations

import argparse
import json
import math
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from depression_rag.evaluation.deidentify import (  # noqa: E402
    SCRUB,
    assert_deidentified,
    deidentify,
)
QUESTIONS = ROOT / "data/derived/questions_alodokter_sample50.jsonl"
ASSIGNMENTS = ROOT / "data/derived/eval_case_assignments.json"
OUTDIR = ROOT / "outputs/evaluation/study_kit/01_KONSELOR"

# --- crisis line -------------------------------------------------------------
# READ FROM configs/chatbot.yaml, never re-typed. These strings are printed into
# all four counselor workbooks, and the counselor reads them beside a chatbot
# that shows `safety.safe_response`. When the two were maintained separately the
# printed sheet drifted: it added 112 (deliberately dropped from the banner) and
# asserted limited hotline hours (unresolved, and not to be stated to
# participants). Emergency services still print FIRST, matching escalation-first.
# To change a number, edit configs/chatbot.yaml — both surfaces follow.
def _escalation_lines() -> tuple[str, str]:
    import yaml

    saf = yaml.safe_load((ROOT / "configs" / "chatbot.yaml").read_text(
        encoding="utf-8"))["safety"]
    esc = saf.get("escalation") or {}
    emergency, crisis = esc.get("emergency", "").strip(), esc.get("crisis", "").strip()
    if not emergency or not crisis:
        raise SystemExit(
            "configs/chatbot.yaml is missing safety.escalation.{emergency,crisis}; "
            "the counselor workbooks derive their crisis lines from it.")
    # The printed sheet must not contradict the banner the counselor sees in the
    # tool. These are the two facts that actually drifted before.
    banner = saf.get("safe_response", "")
    if "119" in banner and "119" not in crisis:
        raise SystemExit("safety.escalation.crisis omits the 119 line that safe_response carries.")
    if "112" in emergency or "112" in crisis:
        raise SystemExit("112 is deliberately absent from safe_response; do not print it.")
    return emergency, crisis


EMERGENCY_LINE, CRISIS_LINE = _escalation_lines()

# --- palette -----------------------------------------------------------------
INK = "1F2933"
MUTED = "5B6875"
LINE = "D5DCE4"
ACCENT = {"DENGAN": "0F6B66", "TANPA": "8A5A12"}
TINT = {"DENGAN": "E8F3F2", "TANPA": "FAF1E2"}
SOFT = "F4F6F8"
WARN_BG = "FDECEC"
WARN_INK = "9B2C2C"

THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# --- answer-sheet option lists (shown verbatim in the dropdowns) -------------
OPTION_LISTS: dict[str, list[str]] = {
    "bantu": [
        "1 – tidak membantu",
        "2 – kurang membantu",
        "3 – cukup membantu",
        "4 – membantu",
        "5 – sangat membantu",
    ],
    "bantu_risiko": [
        "1 – tidak membantu",
        "2 – kurang membantu",
        "3 – cukup membantu",
        "4 – membantu",
        "5 – sangat membantu",
        "Tidak ada tanda bahaya pada kasus ini",
    ],
    "pengaruh": [
        "Tidak memengaruhi",
        "Sedikit memengaruhi",
        "Cukup memengaruhi",
        "Sangat memengaruhi",
    ],
    "yakin": [
        "1 – sangat tidak yakin",
        "2 – kurang yakin",
        "3 – cukup yakin",
        "4 – yakin",
        "5 – sangat yakin",
    ],
    "keliru": ["Tidak", "Ya"],
    "setuju": [
        "1 – sangat tidak setuju",
        "2 – tidak setuju",
        "3 – netral",
        "4 – setuju",
        "5 – sangat setuju",
    ],
}

# --- exit questionnaire (DENGAN workbooks only) ------------------------------
# Numbered 1..N for the counselor. The mapping back to internal codes and source
# constructs lives in the internal notes file, not here.
EXIT_ITEMS: list[tuple[str, str]] = [
    ("head", "Manfaat alat bantu"),
    ("skala", "Alat bantu membantu saya menentukan langkah penanganan dan rujukan yang tepat."),
    ("skala", "Isi ringkasan yang saya terima relevan dan akurat untuk kasus-kasus yang saya "
              "tangani."),
    ("skala", "Dengan alat bantu, jawaban yang saya tulis menjadi lebih baik dibandingkan bila "
              "saya menjawab sendiri."),
    ("skala", "Alat bantu menghemat waktu dan tenaga saya."),
    ("head", "Kemudahan pemakaian"),
    ("skala", "Alat bantu ini mudah digunakan."),
    ("skala", "Isi ringkasannya mudah dipahami — bahasa dan susunannya jelas."),
    ("head", "Sikap terhadap isinya"),
    ("skala", "Saya mempercayai isi ringkasan yang diberikan."),
    ("skala", "Saya tetap memeriksa dan menimbang sendiri isi ringkasan sebelum memakainya."),
    ("head", "Rencana pemakaian"),
    ("skala", "Bila tersedia, saya akan memakai alat seperti ini dalam praktik saya."),
    ("head", "Penilaian menyeluruh"),
    ("angka", "Secara keseluruhan, seberapa berguna alat bantu ini bagi Anda? Isilah angka "
              "0–10  (0 = sama sekali tidak berguna, 10 = sangat berguna)."),
    ("head", "Pertanyaan terbuka"),
    ("teks", "Bagian mana dari alat bantu yang paling membantu Anda?"),
    ("teks", "Bagian mana yang kurang membantu atau perlu diperbaiki?"),
    ("teks", "Hal lain yang ingin Anda sampaikan? (boleh dikosongkan)"),
]


# ================================================================ data ========
def load_cases() -> dict[str, list[dict]]:
    """-> {"counselor_1": [{study_id, question_id, block, condition, question}, ...], ...}"""
    questions = {}
    with QUESTIONS.open(encoding="utf-8") as fh:
        for line in fh:
            rec = json.loads(line)
            questions[rec["question_id"]] = rec
    assign = json.loads(ASSIGNMENTS.read_text(encoding="utf-8"))

    # SCRUB is shared with the psychologist packets, so a target may legitimately
    # sit outside the counselor sample — it just has to exist in the corpus.
    stray = set(SCRUB) - set(questions)
    if stray:
        raise SystemExit(f"SCRUB targets not in the study corpus: {sorted(stray)}")

    out: dict[str, list[dict]] = {"counselor_1": [], "counselor_2": []}
    scrubbed: list[str] = []
    for case in assign["cases"]:
        if not case.get("in_counselor_sample"):
            continue
        qid = case["question_id"]
        text = deidentify(qid, questions[qid]["question"])
        assert_deidentified(text, f'{case["study_id"]} (counselor workbook)')
        if qid in SCRUB:
            scrubbed.append(case["study_id"])
        for counselor, cond in case["counselor_condition"].items():
            out[counselor].append(
                {
                    "study_id": case["study_id"],
                    "question_id": qid,
                    "block": str(case["block"]),
                    "condition": cond,
                    "question": text,
                }
            )
    for lst in out.values():
        lst.sort(key=lambda c: c["study_id"])
    print(f"  de-identified: {sorted(set(scrubbed))}")
    return out


# ============================================================== layout ========
def _put(ws, row: int, col: int, value=None, *, font=None, fill=None, align=None,
         border=None, number_format=None):
    cell = ws.cell(row=row, column=col, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    if number_format:
        cell.number_format = number_format
    return cell


def _band(ws, row: int, first_col: int, last_col: int, fill: str) -> None:
    """Paint every cell of a merged band so Excel renders the fill across it."""
    for c in range(first_col, last_col + 1):
        _put(ws, row, c, fill=fill)


def _lines(text, width: int) -> int:
    return sum(max(1, math.ceil(len(part) / width)) for part in str(text).split("\n"))


def _fit_to_width(ws, *, landscape: bool, print_area: str, title_rows: str | None = None) -> None:
    """Keep every column on one printed page — counselors do print these."""
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_area = print_area
    if title_rows:
        ws.print_title_rows = title_rows


# --------------------------------------------------------- instructions ------
def build_petunjuk(wb: Workbook, *, counselor: int, cond: str, n_cases: int,
                   stage: int, other_stage_file: str) -> None:
    """Sheet 1 — everything the counselor needs, read once, top to bottom."""
    ws = wb.create_sheet("Petunjuk")
    accent, tint = ACCENT[cond], TINT[cond]
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 76
    MERGED_W = 110

    f_title = Font(name="Calibri", size=18, bold=True, color=accent)
    f_sub = Font(name="Calibri", size=11, color=MUTED)
    f_head = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    f_p = Font(name="Calibri", size=10.5, color=INK)
    f_b = Font(name="Calibri", size=10.5, bold=True, color=INK)
    f_warn = Font(name="Calibri", size=10.5, bold=True, color=WARN_INK)
    f_lbl = Font(name="Calibri", size=10.5, bold=True, color=accent)
    f_th = Font(name="Calibri", size=10.5, bold=True, color="FFFFFF")
    top = Alignment(vertical="top", wrap_text=True)
    top_i = Alignment(vertical="top", wrap_text=True, indent=1)
    mid = Alignment(vertical="center")

    r = 1

    def para(text, font=f_p, fill=None, indent=False, extra=0):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        if fill:
            _band(ws, r, 2, 3, fill)
        _put(ws, r, 2, text, font=font, fill=fill, align=top_i if indent else top)
        ws.row_dimensions[r].height = _lines(text, MERGED_W) * 13.5 + 5 + extra
        r += 1

    def head(text):
        nonlocal r
        ws.row_dimensions[r].height = 8
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        _band(ws, r, 2, 3, accent)
        _put(ws, r, 2, "  " + text, font=f_head, fill=accent, align=mid)
        ws.row_dimensions[r].height = 23
        r += 1

    def kv(label, value, label_font=f_lbl, fill=None, box=False):
        nonlocal r
        b = BOX if box else None
        _put(ws, r, 2, label, font=label_font, fill=fill, align=top, border=b)
        _put(ws, r, 3, value, font=f_p, fill=fill, align=top, border=b)
        ws.row_dimensions[r].height = max(_lines(label, 34), _lines(value, 74)) * 13.5 + 6
        r += 1

    def gap(h=6):
        nonlocal r
        ws.row_dimensions[r].height = h
        r += 1

    cond_label = "DENGAN alat bantu" if cond == "DENGAN" else "TANPA alat bantu"

    # ---- header -------------------------------------------------------------
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    _put(ws, r, 2, f"LEMBAR KERJA KONSELOR — {cond_label.upper()}", font=f_title, align=mid)
    ws.row_dimensions[r].height = 27
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    _put(ws, r, 2,
         f"Konselor {counselor}   ·   Tahap {stage} dari 2   ·   {n_cases} kasus   ·   "
         f"perkiraan waktu {'2–3' if cond == 'DENGAN' else '2'} jam",
         font=f_sub, align=mid)
    ws.row_dimensions[r].height = 18
    r += 1
    gap(4)
    para("Bacalah halaman ini satu kali sampai habis sebelum mulai. Semua yang Anda perlukan "
         "ada di sini — tugasnya, cara menulis jawaban, aturannya, dan arti setiap kolom yang "
         "harus Anda isi — sehingga Anda tidak perlu bolak-balik saat mengerjakan.",
         font=f_b, fill=tint, extra=6)

    # ---- identity -----------------------------------------------------------
    head("Identitas berkas ini")
    kv("Kode konselor", f"Konselor {counselor}", fill=SOFT, box=True)
    kv("Kondisi pengerjaan", cond_label, fill=SOFT, box=True)
    kv("Tanggal mulai", None, box=True)
    kv("Tanggal selesai", None, box=True)
    gap()
    para("Jangan menuliskan nama Anda di berkas ini. Anda dikenali hanya dengan kode di atas, "
         "dan jawaban Anda dinilai tanpa diketahui siapa penulisnya.", indent=True)

    # ---- 1. the task --------------------------------------------------------
    head("1. Apa yang Anda kerjakan")
    para(f"Di sheet «Jawaban» ada {n_cases} pertanyaan nyata dari masyarakat mengenai masalah "
         "depresi dan kesehatan jiwa. Untuk setiap pertanyaan, tulislah satu pesan balasan "
         "yang ditujukan langsung kepada penanya — seolah Anda sedang membalas orang tersebut.",
         indent=True)
    para("Yang Anda hasilkan adalah pesan untuk penanya: bukan catatan klinis, bukan ringkasan "
         "pedoman, bukan daftar poin berjudul.", indent=True)
    if cond == "DENGAN":
        para("Pada tahap ini Anda memakai alat bantu. Alat itu memberi Anda ringkasan "
             "pertimbangan untuk setiap kasus; jawaban akhir untuk penanya tetap Anda yang "
             "menulis. Cara memakainya ada di bagian 4.", indent=True)
    else:
        para("Pada tahap ini Anda menjawab murni dari pengetahuan dan pengalaman Anda sendiri, "
             "tanpa alat bantu apa pun.", indent=True)

    # ---- 2. how to answer ---------------------------------------------------
    head("2. Cara menulis jawaban")
    para("Tulislah sebagai paragraf yang mengalir dan hangat: satu balasan utuh sepanjang "
         "kira-kira 150–350 kata, dalam Bahasa Indonesia yang jelas dan tidak menghakimi. "
         "Isinya sebaiknya mencakup lima hal berikut, secara berurutan:", indent=True)
    for i, (t, d) in enumerate(
        [
            ("Akui perasaannya lebih dulu",
             "sapa penanya dan validasi apa yang sedang ia rasakan."),
            ("Jelaskan secara awam",
             "kemungkinan apa yang sedang terjadi. Jangan menegakkan diagnosis pasti dan "
             "jangan meresepkan obat."),
            ("Beri langkah konkret",
             "hal yang bisa penanya lakukan sendiri mulai hari ini."),
            # Referenced BY NAME, not by number. The danger-signs section is
            # numbered 5 in the DENGAN workbook but 4 in TANPA (the tool section
            # only exists in DENGAN, see `n` below), so the old "lihat bagian 5"
            # sent TANPA readers to «Arti setiap kolom» — a column glossary —
            # on the one cross-reference that carries safety weight. The number
            # also collided with this list's own 1–5 items, where 5 is
            # "Tutup dengan menguatkan", the opposite of escalating.
            ("Arahkan ke bantuan profesional",
             "kapan dan ke mana harus mencari bantuan. Bila ada tanda bahaya, bagian inilah "
             "yang didahulukan — lihat bagian «Bila ada tanda bahaya» di bawah."),
            ("Tutup dengan menguatkan",
             "kalimat penutup yang memberi harapan."),
        ],
        start=1,
    ):
        kv(f"{i}.  {t}", d)
    gap()
    para("Lima hal di atas adalah isi yang perlu tercakup, bukan format tampilan. Tulislah "
         "dengan gaya Anda sendiri sebagai satu balasan yang utuh — jangan memberi judul "
         "bagian atau membuat daftar bernomor di dalam jawaban.", indent=True)
    para("Pakailah cara menulis yang sama persis pada kedua tahap, agar jawaban Anda dapat "
         "dibandingkan secara adil.", font=f_b, fill=tint, extra=4)

    # ---- 3. rules -----------------------------------------------------------
    head("3. Aturan yang harus dipatuhi")
    rules = [
        "Kerjakan sendiri. Jangan mendiskusikan pertanyaan atau jawaban dengan konselor lain.",
        "Jawab semua pertanyaan sesuai urutan, termasuk yang terasa sulit.",
        "Setelah pindah ke pertanyaan berikutnya, jangan kembali mengubah jawaban sebelumnya.",
        "Jangan mencantumkan identitas atau data pribadi penanya di dalam jawaban.",
    ]
    if cond == "DENGAN":
        rules += [
            "Selain alat bantu yang diberikan, jangan memakai internet, mesin pencari, buku "
            "pedoman, atau aplikasi AI lain.",
            "Jangan menyebutkan di dalam jawaban bahwa Anda memakai alat bantu.",
        ]
    else:
        rules += [
            "Jangan memakai internet, mesin pencari, aplikasi AI atau chatbot apa pun, buku "
            "pedoman, maupun bertanya kepada rekan.",
            "Sebaiknya selesaikan seluruh tahap ini dalam satu atau dua sesi saja.",
        ]
    for rule in rules:
        para("•   " + rule, indent=True)

    # ---- 4. the tool --------------------------------------------------------
    n = 4
    if cond == "DENGAN":
        head("4. Cara memakai alat bantu")
        kv("Tautan alat bantu Anda", "[TEMPEL TAUTAN DI SINI]", fill=SOFT, box=True)
        gap()
        # RESOLVED 2026-08-08. This said "salin pertanyaan penanya ... ke alat
        # bantu" while Pengarahan_Konselor.pptx slide 6 said "Mohon pertanyaannya
        # tidak disalin-tempel". Whichever a counselor followed was arbitrary, so
        # the primary input to the retriever was uncontrolled. Paraphrase wins:
        # it is what the briefing deck instructs, and it is why the study-question
        # picker was removed from the UI on 2026-07-27 (chatbot_app.py:197-198) —
        # the tool must see the counselor's own framing.
        para("Untuk setiap kasus: tuliskan ulang kasus penanya dengan kata-kata Anda sendiri "
             "ke alat bantu — JANGAN disalin-tempel dari kolom «Pertanyaan dari penanya». "
             "Baca ringkasan yang muncul, lalu tulis sendiri jawaban akhir untuk penanya di "
             "kolom «Jawaban Anda».", indent=True)
        para("Alat bantu ini tidak mengingat pesan sebelumnya: setiap pertanyaan diproses "
             "terpisah, jadi sertakan kembali konteks kasus pada tiap pertanyaan.", indent=True)
        para("Tautan bersifat pribadi untuk Anda — jangan diteruskan.", indent=True)
        # The workbook carried NO logging disclosure at all, and named no processor —
        # the briefing deck's one vague line was the only mention, and a counselor
        # working through cases has the workbook open, not the deck. Storage and
        # transit are stated together because they are one question — where does the
        # text go — and because together they are what justifies the rule that
        # follows, which previously stood without a reason.
        para("KE MANA TEKS ANDA PERGI. Dua hal yang perlu Anda ketahui sebelum mulai:",
             font=f_b, fill=tint, extra=4)
        para("•   TERSIMPAN. Teks pertanyaan Anda, ringkasan yang dihasilkan, dan waktunya "
             "dicatat apa adanya untuk keperluan analisis penelitian. Itulah datanya untuk "
             "kondisi DENGAN alat bantu.", indent=True)
        para("•   MELEWATI PIHAK KETIGA. Alat bantu ini dibuka lewat tautan publik sementara "
             "(gradio.live) yang dikelola pihak ketiga. Teks yang Anda kirim dan ringkasan "
             "yang muncul melewati server tersebut sebelum sampai ke komputer peneliti.",
             indent=True)
        para("Karena itu: gunakan HANYA kasus studi ini. Jangan mengetikkan nama, kontak, "
             "atau data klien nyata.", font=f_warn, fill=WARN_BG, extra=4)
        para("Alat bantu boleh Anda ikuti, boleh juga Anda abaikan sebagian atau seluruhnya. "
             "Penilaian Anda sendiri tetap yang menentukan.", indent=True)
        gap()
        para("Agar penilaian tetap objektif, jawaban akhir Anda TIDAK BOLEH memuat:",
             font=f_warn, fill=WARN_BG, extra=4)
        for bad in [
            "judul bagian yang dipakai alat bantu (misalnya «INFORMASI PEDOMAN» atau "
            "«PERTIMBANGAN KLINIS»);",
            "nomor rujukan dalam kurung siku seperti [1] atau [2];",
            "kode atau nama bagian pedoman, misalnya «MI.7»;",
            "kalimat yang disalin mentah-mentah dari ringkasan.",
        ]:
            para("•   " + bad, indent=True)
        gap()
        para("Jawaban akhir harus terbaca sebagai balasan alami kepada penanya, dengan gaya dan "
             "format yang sama seperti bila Anda menjawab tanpa alat bantu.", indent=True)
        n = 5

    # ---- safety -------------------------------------------------------------
    head(f"{n}. Bila ada tanda bahaya")
    para("Sebagian pertanyaan memuat tanda bahaya: keinginan bunuh diri, rencana atau usaha "
         "menyakiti diri, atau keadaan gawat darurat lain. Bila Anda menemukannya:", indent=True)
    for s in [
        "dahulukan keselamatan — tanggapi tanda bahaya itu secara langsung dan empatik, jangan "
        "dilewati;",
        "sarankan bantuan profesional segera, dan sebutkan jalur bantuan darurat di bawah ini;",
        "jangan menangani krisis sendirian lewat pesan — arahkan penanya kepada tenaga "
        "kesehatan.",
    ]:
        para("•   " + s, indent=True)
    gap()
    kv("Gawat darurat (pertama)", EMERGENCY_LINE, fill=WARN_BG, box=True)
    kv("Dukungan psikologis", CRISIS_LINE, fill=WARN_BG, box=True)
    gap()
    para("Gunakan nomor yang sama pada semua jawaban yang membutuhkannya.", indent=True)

    # ---- columns ------------------------------------------------------------
    head(f"{n + 1}. Arti setiap kolom di sheet «Jawaban»")
    para("Isilah baris demi baris, dari atas ke bawah. Sel berlatar putih adalah sel yang harus "
         "Anda isi; sel berlatar abu-abu sudah terisi atau terhitung sendiri — biarkan saja.",
         indent=True)
    gap()
    _put(ws, r, 2, "Kolom", font=f_th, fill=MUTED, align=top, border=BOX)
    _put(ws, r, 3, "Cara mengisi", font=f_th, fill=MUTED, align=top, border=BOX)
    ws.row_dimensions[r].height = 18
    r += 1

    cols: list[tuple[str, str]] = [
        ("Kode kasus",
         "Sudah terisi. Jangan diubah — kode inilah yang menghubungkan jawaban Anda dengan "
         "pertanyaannya."),
        ("Pertanyaan dari penanya",
         "Sudah terisi. Bila teksnya panjang, klik selnya dan bacalah pada baris rumus di atas, "
         "atau tarik batas barisnya ke bawah."),
        ("Jawaban Anda",
         "Ketik balasan Anda di sini, 150–350 kata. Untuk berganti baris di dalam sel, tekan "
         "Alt+Enter."),
        ("Jumlah kata",
         "Terhitung sendiri. Selnya berubah merah bila jawaban Anda kurang dari 150 atau lebih "
         "dari 350 kata."),
        ("Waktu mulai  /  Waktu selesai",
         "Tulis jam saat Anda mulai dan selesai mengerjakan kasus itu, format 24 jam — "
         "misalnya 09:15."),
        ("Lama (menit)", "Terhitung sendiri dari kedua jam di atas."),
    ]
    if cond == "DENGAN":
        cols += [
            ("Seberapa membantu ringkasan untuk kasus ini?",
             "Pilih 1–5 dari daftar. Isilah segera setelah jawaban Anda selesai, selagi kesan "
             "Anda masih jelas."),
            ("Seberapa membantu ringkasan untuk mengenali tanda bahaya?",
             "Pilih 1–5 dari daftar. Bila menurut Anda kasus ini memang tidak memuat tanda "
             "bahaya, pilih pilihan yang terakhir."),
            ("Seberapa besar ringkasan memengaruhi jawaban akhir Anda?",
             "Pilih satu. Ini soal apakah isi jawaban Anda berubah karena ringkasan tersebut — "
             "bukan soal bagus atau tidaknya ringkasan itu."),
            ("Seberapa yakin Anda pada jawaban ini?",
             "Pilih 1–5 dari daftar: seberapa yakin Anda terhadap jawaban yang baru saja Anda "
             "tulis."),
            ("Ada isi ringkasan yang keliru atau berpotensi membahayakan?",
             "Pilih «Tidak» atau «Ya». Pertanyaan ini penting — jawablah apa adanya."),
            ("Bila «Ya», jelaskan singkat",
             "Cukup satu–dua kalimat: bagian mana yang keliru dan mengapa."),
        ]
    else:
        cols += [
            ("Seberapa yakin Anda pada jawaban ini?",
             "Pilih 1–5 dari daftar, segera setelah jawaban Anda selesai: seberapa yakin Anda "
             "terhadap jawaban yang baru saja Anda tulis."),
        ]
    for label, how in cols:
        kv(label, how, label_font=f_b, box=True)

    # ---- exit questionnaire -------------------------------------------------
    if cond == "DENGAN":
        head(f"{n + 2}. Kuesioner akhir")
        para(f"Setelah seluruh {n_cases} kasus di sheet «Jawaban» selesai, bukalah sheet "
             "«Kuesioner Akhir» dan isilah satu kali. Isinya adalah pendapat Anda secara "
             "keseluruhan tentang alat bantu ini. Tidak ada jawaban benar atau salah — "
             "penilaian yang jujur, termasuk yang negatif, justru yang paling berguna.",
             indent=True)
        last = n + 3
    else:
        last = n + 2

    # ---- checklist ----------------------------------------------------------
    head(f"{last}. Sebelum berkas ini dikirim kembali")
    checks = [
        f"Semua {n_cases} baris di sheet «Jawaban» sudah terisi.",
        "Setiap jawaban berada di rentang 150–350 kata (kolom «Jumlah kata» tidak merah).",
        "Jam mulai dan jam selesai terisi untuk setiap kasus.",
    ]
    if cond == "DENGAN":
        checks += [
            "Kolom penilaian di sebelah kanan terisi untuk setiap kasus.",
            "Sheet «Kuesioner Akhir» sudah diisi satu kali.",
            "Tidak ada judul bagian, nomor rujukan [1], atau kode pedoman yang tersalin ke "
            "dalam jawaban.",
        ]
    else:
        checks.append("Kolom keyakinan terisi untuk setiap kasus.")
    checks.append("Nama Anda tidak tertulis di mana pun dalam berkas ini.")
    for c in checks:
        para("☐   " + c, indent=True)
    gap()
    para(f"Setelah tahap ini selesai, lanjutkan ke berkas «{other_stage_file}». Jangan "
         "membukanya lebih dulu."
         if stage == 1 else
         "Ini tahap terakhir. Setelah selesai, simpan berkas ini dan kirimkan kedua berkas Anda "
         "kembali kepada peneliti.",
         font=f_b, fill=tint, extra=4)
    gap()
    # "jangan menebak" was a bare prohibition aimed at a professional participant,
    # so it read as distrust — as if we expected them to guess. Same instruction,
    # framed as an invitation to ask rather than an accusation.
    kv("Ada yang kurang jelas?",
       "Silakan hubungi peneliti sebelum melanjutkan — lebih baik bertanya "
       "daripada berasumsi.", label_font=f_b)

    _fit_to_width(ws, landscape=False, print_area=f"B1:C{r - 1}")


# --------------------------------------------------------- answer sheet ------
HEAD_COMMON = [
    ("No", 5), ("Kode kasus", 11),
    ("Pertanyaan dari penanya", 62), ("Jawaban Anda  (150–350 kata)", 62),
    ("Jumlah kata", 9), ("Waktu mulai", 11), ("Waktu selesai", 11), ("Lama (menit)", 10),
]
HEAD_DENGAN = HEAD_COMMON + [
    ("Seberapa membantu ringkasan untuk kasus ini?", 24),
    ("Seberapa membantu ringkasan untuk mengenali tanda bahaya?", 26),
    ("Seberapa besar ringkasan memengaruhi jawaban akhir Anda?", 25),
    ("Seberapa yakin Anda pada jawaban ini?", 23),
    ("Ada isi ringkasan yang keliru atau berpotensi membahayakan?", 21),
    ("Bila «Ya», jelaskan singkat", 34),
]
HEAD_TANPA = HEAD_COMMON + [("Seberapa yakin Anda pada jawaban ini?", 23)]

DV_COLS = {
    "DENGAN": {"I": "bantu", "J": "bantu_risiko", "K": "pengaruh", "L": "yakin", "M": "keliru"},
    "TANPA": {"I": "yakin"},
}


def build_jawaban(wb: Workbook, cases: list[dict], *, cond: str, counselor: int,
                  option_ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("Jawaban")
    accent, tint = ACCENT[cond], TINT[cond]
    ws.sheet_view.showGridLines = False

    f_h = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    f_q = Font(name="Calibri", size=10, color=INK)
    f_meta = Font(name="Calibri", size=10, color=MUTED)
    top = Alignment(vertical="top", wrap_text=True)
    ctr = Alignment(vertical="center", horizontal="center", wrap_text=True)

    headers = HEAD_DENGAN if cond == "DENGAN" else HEAD_TANPA
    ncol = len(headers)

    banner = ("Isilah dari atas ke bawah. Tulis jawaban Anda di kolom «Jawaban Anda», lalu isi "
              "kolom penilaian di sebelah kanan sebelum lanjut ke baris berikutnya."
              if cond == "DENGAN" else
              "Isilah dari atas ke bawah. Tulis jawaban Anda di kolom «Jawaban Anda», lalu isi "
              "kolom keyakinan sebelum lanjut ke baris berikutnya.")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    _band(ws, 1, 1, ncol, tint)
    _put(ws, 1, 1,
         f"KONSELOR {counselor} — {cond} ALAT BANTU.   {banner}",
         font=Font(name="Calibri", size=10.5, bold=True, color=INK), fill=tint,
         align=Alignment(vertical="center", indent=1))
    ws.row_dimensions[1].height = 20

    for i, (title, width) in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        _put(ws, 2, i, title, font=f_h, fill=accent, align=ctr, border=BOX)
    ws.row_dimensions[2].height = 60
    ws.freeze_panes = "C3"

    for idx, case in enumerate(cases, start=1):
        row = idx + 2
        _put(ws, row, 1, idx, font=f_meta, align=ctr, fill=SOFT, border=BOX)
        _put(ws, row, 2, case["study_id"], font=f_meta, align=ctr, fill=SOFT, border=BOX)
        _put(ws, row, 3, case["question"], font=f_q, align=top, fill=SOFT, border=BOX)
        _put(ws, row, 4, align=top, fill="FFFFFF", border=BOX)
        _put(ws, row, 5,
             f'=IF($D{row}="","",'
             f'LEN(TRIM(SUBSTITUTE($D{row},CHAR(10)," ")))'
             f'-LEN(SUBSTITUTE(TRIM(SUBSTITUTE($D{row},CHAR(10)," "))," ",""))+1)',
             font=f_meta, align=ctr, fill=SOFT, border=BOX)
        _put(ws, row, 6, align=ctr, fill="FFFFFF", border=BOX, number_format="hh:mm")
        _put(ws, row, 7, align=ctr, fill="FFFFFF", border=BOX, number_format="hh:mm")
        _put(ws, row, 8,
             f'=IF(OR($F{row}="",$G{row}=""),"",ROUND(MOD($G{row}-$F{row},1)*1440,0))',
             font=f_meta, align=ctr, fill=SOFT, border=BOX, number_format="0")
        for c in range(9, ncol + 1):
            _put(ws, row, c, align=top, fill="FFFFFF", border=BOX)
        ws.row_dimensions[row].height = min(300, max(96, _lines(case["question"], 60) * 13.0))

    last = len(cases) + 2
    ws.conditional_formatting.add(
        f"E3:E{last}",
        FormulaRule(formula=['AND($E3<>"",OR($E3<150,$E3>350))'],
                    fill=PatternFill("solid", fgColor=WARN_BG),
                    font=Font(name="Calibri", size=10, bold=True, color=WARN_INK)),
    )

    _fit_to_width(ws, landscape=True,
                  print_area=f"A1:{get_column_letter(ncol)}{last}", title_rows="1:2")

    for col, key in DV_COLS[cond].items():
        dv = DataValidation(type="list", formula1=option_ranges[key], allow_blank=True,
                            showDropDown=False)
        dv.errorTitle = "Pilihan tidak dikenal"
        dv.error = "Pilihlah salah satu dari daftar yang tersedia."
        ws.add_data_validation(dv)
        dv.add(f"{col}3:{col}{last}")


# ------------------------------------------------------------ exit sheet -----
def build_kuesioner(wb: Workbook, *, counselor: int, option_ranges: dict[str, str]) -> None:
    ws = wb.create_sheet("Kuesioner Akhir")
    accent, tint = ACCENT["DENGAN"], TINT["DENGAN"]
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 5
    ws.column_dimensions["C"].width = 84
    ws.column_dimensions["D"].width = 28

    f_p = Font(name="Calibri", size=10.5, color=INK)
    f_head = Font(name="Calibri", size=10.5, bold=True, color="FFFFFF")
    top = Alignment(vertical="top", wrap_text=True)
    ctr = Alignment(vertical="center", horizontal="center")
    mid = Alignment(vertical="center")

    r = 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    _put(ws, r, 2, "KUESIONER AKHIR",
         font=Font(name="Calibri", size=16, bold=True, color=accent), align=mid)
    ws.row_dimensions[r].height = 25
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    _put(ws, r, 2, f"Konselor {counselor}   ·   diisi SATU KALI, setelah semua kasus di sheet "
                   "«Jawaban» selesai",
         font=Font(name="Calibri", size=10.5, color=MUTED), align=mid)
    ws.row_dimensions[r].height = 18
    r += 2

    intro = ("Pertanyaan berikut menanyakan pendapat Anda secara keseluruhan tentang alat bantu "
             "yang Anda pakai — bukan tentang satu kasus tertentu. Untuk pernyataan bernomor, "
             "pilihlah satu jawaban dari daftar. Tidak ada jawaban benar atau salah, dan "
             "penilaian yang jujur — termasuk yang negatif — justru yang paling berguna.")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    _band(ws, r, 2, 4, tint)
    _put(ws, r, 2, intro, font=f_p, fill=tint, align=top)
    ws.row_dimensions[r].height = _lines(intro, 112) * 13.5 + 6
    r += 2

    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    _band(ws, r, 2, 3, SOFT)
    _put(ws, r, 2, "Tanggal pengisian",
         font=Font(name="Calibri", size=10.5, bold=True, color=accent),
         align=Alignment(vertical="center", indent=1), border=BOX, fill=SOFT)
    _put(ws, r, 3, border=BOX)
    _put(ws, r, 4, align=top, border=BOX)
    ws.row_dimensions[r].height = 20
    r += 2

    _put(ws, r, 2, "#", font=f_head, fill=MUTED, align=ctr, border=BOX)
    _put(ws, r, 3, "Pernyataan / pertanyaan", font=f_head, fill=MUTED, align=top, border=BOX)
    _put(ws, r, 4, "Jawaban Anda", font=f_head, fill=MUTED, align=ctr, border=BOX)
    ws.row_dimensions[r].height = 20
    r += 1

    dv_skala = DataValidation(type="list", formula1=option_ranges["setuju"],
                              allow_blank=True, showDropDown=False)
    dv_skala.errorTitle = "Pilihan tidak dikenal"
    dv_skala.error = "Pilihlah salah satu dari daftar yang tersedia."
    ws.add_data_validation(dv_skala)
    dv_angka = DataValidation(type="whole", operator="between", formula1=0, formula2=10,
                              allow_blank=True)
    dv_angka.errorTitle = "Angka di luar rentang"
    dv_angka.error = "Isilah angka bulat dari 0 sampai 10."
    ws.add_data_validation(dv_angka)

    num = 0
    for kind, text in EXIT_ITEMS:
        if kind == "head":
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
            _band(ws, r, 2, 4, accent)
            _put(ws, r, 2, "  " + text, font=f_head, fill=accent, align=mid)
            ws.row_dimensions[r].height = 20
            r += 1
            continue
        num += 1
        _put(ws, r, 2, num, font=f_p, align=ctr, border=BOX, fill=SOFT)
        _put(ws, r, 3, text, font=f_p, align=top, border=BOX, fill=SOFT)
        _put(ws, r, 4, align=top, border=BOX)
        if kind == "skala":
            dv_skala.add(f"D{r}")
            ws.row_dimensions[r].height = max(20, _lines(text, 82) * 13.5 + 6)
        elif kind == "angka":
            dv_angka.add(f"D{r}")
            ws.row_dimensions[r].height = max(20, _lines(text, 82) * 13.5 + 6)
        else:
            ws.row_dimensions[r].height = 64
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    _band(ws, r, 2, 4, tint)
    _put(ws, r, 2, "Terima kasih. Simpan berkas ini, lalu kirimkan kembali kepada peneliti.",
         font=Font(name="Calibri", size=10.5, bold=True, color=INK), fill=tint, align=mid)
    ws.row_dimensions[r].height = 22

    _fit_to_width(ws, landscape=False, print_area=f"B1:D{r}")


# -------------------------------------------------------- option sheet -------
def option_ranges() -> dict[str, str]:
    return {
        key: f"_pilihan!${get_column_letter(i)}$2:${get_column_letter(i)}${len(v) + 1}"
        for i, (key, v) in enumerate(OPTION_LISTS.items(), start=1)
    }


def build_options(wb: Workbook) -> None:
    """Hidden sheet holding the dropdown lists — created last, stays last."""
    ws = wb.create_sheet("_pilihan")
    for i, (key, values) in enumerate(OPTION_LISTS.items(), start=1):
        ws.cell(row=1, column=i, value=key)
        for j, v in enumerate(values, start=2):
            ws.cell(row=j, column=i, value=v)
    ws.sheet_state = "hidden"


# ================================================================ build =======
FNAME = {
    ("DENGAN", 1): "Konselor_1_DENGAN_alat_bantu.xlsx",
    ("TANPA", 1): "Konselor_1_TANPA_alat_bantu.xlsx",
    ("DENGAN", 2): "Konselor_2_DENGAN_alat_bantu.xlsx",
    ("TANPA", 2): "Konselor_2_TANPA_alat_bantu.xlsx",
}


def build_workbook(cases: list[dict], *, counselor: int, cond: str, stage: int,
                   other_stage_file: str, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    opts = option_ranges()
    build_petunjuk(wb, counselor=counselor, cond=cond, n_cases=len(cases), stage=stage,
                   other_stage_file=other_stage_file)
    build_jawaban(wb, cases, cond=cond, counselor=counselor, option_ranges=opts)
    if cond == "DENGAN":
        build_kuesioner(wb, counselor=counselor, option_ranges=opts)
    build_options(wb)
    wb.active = 0
    wb.save(path)
    visible = [s for s in wb.sheetnames if s != "_pilihan"]
    print(f"  wrote {path.relative_to(ROOT)}  ({len(cases)} kasus, sheets={visible})")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--outdir", default=str(OUTDIR))
    ap.add_argument("--archive-old", action="store_true",
                    help="move the superseded combined workbooks into _ARSIP_versi_lama/")
    args = ap.parse_args(argv)
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)

    per_counselor = load_cases()
    for counselor in (1, 2):
        cases = per_counselor[f"counselor_{counselor}"]
        chatbot = [c for c in cases if c["condition"] == "chatbot"]
        plain = [c for c in cases if c["condition"] == "no_chatbot"]
        if not len(chatbot) == len(plain) == 12:
            raise SystemExit(f"counselor {counselor}: expected 12+12, got "
                             f"{len(chatbot)}+{len(plain)}")
        # the counterbalanced crossover: whoever holds block 1 works it first
        first = "DENGAN" if chatbot[0]["block"] == "1" else "TANPA"
        for cond, subset in (("DENGAN", chatbot), ("TANPA", plain)):
            other = FNAME[("TANPA" if cond == "DENGAN" else "DENGAN", counselor)]
            build_workbook(subset, counselor=counselor, cond=cond,
                           stage=1 if cond == first else 2, other_stage_file=other,
                           path=outdir / FNAME[(cond, counselor)])

    if args.archive_old:
        arsip = outdir.parent / "_ARSIP_versi_lama"
        arsip.mkdir(exist_ok=True)
        for old in sorted(outdir.glob("Konselor_*_LEMBAR_JAWABAN.xlsx")):
            shutil.move(str(old), str(arsip / old.name))
            print(f"  archived {old.name} -> {arsip.relative_to(ROOT)}/")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
